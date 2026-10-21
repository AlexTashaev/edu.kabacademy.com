"""«Активность студентов КК1» — еженедельная выгрузка по курсу 236 в Google-таблицу.

Курс 236 «Классическая каббала. Часть первая (Осень 2026)». По каждому студенту:
группа, ссылка на профиль, по каждому уроку — лекция пройдена, был на вебинаре
в Zoom, открывал архив вебинара; сколько раз входил на сайт с начала обучения
и когда был на сайте последний раз.

Еженедельный запуск — cron на web-18 (пользователь edu.kabacademy.com, под уже
стоящим в crontab CRON_TZ=Asia/Jerusalem), файлы в /sites/edu.kabacademy.com/public/kab-tools/kk1_activity/:
    0 11 * * 5 cd <папка> && PYTHONIOENCODING=utf-8 /usr/bin/python3 kk1_activity.py --local-sql ./moodle_sql.sh --push ./config.json >> run.log 2>&1

Разовый запуск с машины владельца (ssh-алиас web-18, нужен openpyxl для -o):
    python kk1_activity.py -o "<папка вне репо>" [--push "<config.json>"]

-o          — куда положить xlsx-копию (в ней ПДн: только вне репозитория); без -o не пишется;
--push      — config.json вне репо: {"url": "<exec-URL Apps Script>", "secret": "..."};
              выгрузка ложится в таблицу новой вкладкой с датой (см. kk1_activity.gs);
--local-sql — SQL через локальную обёртку moodle_sql.sh (на самом web-18), а не через ssh.

Уроки попадают в выгрузку, когда вебинар урока уже начался (по mdl_zoom.start_time).

Откуда что берётся:
- студенты — роль student в курсе без ролей сотрудников, admin id 2 исключён
  (как в webinars_report.py);
- группа — группы курса, где состоит студент;
- лекция пройдена — отметка о выполнении лекции (completion: дочитана до конца);
- Zoom — студент есть в отчёте Zoom (mdl_zoom_meeting_participants, к пользователю
  его привязывает сам mod_zoom по email и похожести имени);
- архив — открывал страницу «Запись вебинара N (архив)»: ролики на YouTube/Drive,
  их просмотр Moodle не видит;
- заходы — события входа на сайт (user_loggedin) с даты начала курса; вход в
  мобильное приложение события не даёт;
- последнее посещение — mdl_user.lastaccess (любое действие, и из приложения тоже).
"""
import argparse
import json
import re
import subprocess
import time
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

COURSE = 236
SITE = "https://edu.kabacademy.com"
TZ = ZoneInfo("Asia/Jerusalem")
ADMINS = {2}  # аккаунт администратора записан студентом — не считать
STAFF_ROLES = {"teacher", "editingteacher", "manager", "forummoderator", "moderator_assist", "koordinator"}

# урок -> (zoom id, cmid «Запись вебинара N (архив)», cmid лекции)
LESSONS = {
    1: (435, 13440, 13459), 2: (408, 13481, 13466), 3: (407, 13490, 13492),
    4: (406, 13507, 13509), 5: (405, 13523, 13525), 6: (404, 13539, 13541),
    7: (403, 13554, 13556), 8: (402, 13569, 13571), 9: (401, 13584, 13586),
    10: (400, 13599, 13601), 11: (399, 13614, 13616), 12: (398, 13629, 13631),
    13: (397, 13644, 13646), 14: (396, 13659, 13661), 15: (395, 13674, 13676),
}
SQL_CMD = ["ssh", "web-18", '/tmp/kab_moodle_sql2.sh "$(cat)"']
LOCAL_SQL = None  # путь к moodle_sql.sh, когда скрипт работает на самом web-18


def ids(xs):
    return ",".join(str(x) for x in xs)


ZOOMS = ids(v[0] for v in LESSONS.values())
ARCHIVES = ids(v[1] for v in LESSONS.values())
LECTURES = ids(v[2] for v in LESSONS.values())

QUERIES = {
    "course": f"SELECT startdate FROM mdl_course WHERE id = {COURSE}",
    "users": f"""
SELECT u.id, u.firstname, u.lastname, u.email, u.lastaccess,
 (SELECT GROUP_CONCAT(DISTINCT r.shortname) FROM mdl_role_assignments ra
   JOIN mdl_role r ON r.id = ra.roleid
   JOIN mdl_context cx ON cx.id = ra.contextid AND cx.contextlevel = 50 AND cx.instanceid = {COURSE}
  WHERE ra.userid = u.id) roles,
 (SELECT GROUP_CONCAT(g.name ORDER BY g.name + 0, g.name SEPARATOR ', ') FROM mdl_groups_members gm
   JOIN mdl_groups g ON g.id = gm.groupid AND g.courseid = {COURSE}
  WHERE gm.userid = u.id) grp
FROM mdl_user u
JOIN mdl_user_enrolments ue ON ue.userid = u.id
JOIN mdl_enrol e ON e.id = ue.enrolid AND e.courseid = {COURSE}
WHERE u.deleted = 0
GROUP BY u.id""",
    "zooms": f"SELECT id, start_time FROM mdl_zoom WHERE id IN ({ZOOMS})",
    "zoom": f"""
SELECT DISTINCT d.zoomid, p.userid
FROM mdl_zoom_meeting_participants p
JOIN mdl_zoom_meeting_details d ON d.id = p.detailsid
WHERE d.zoomid IN ({ZOOMS}) AND p.userid IS NOT NULL""",
    "lectures": f"""
SELECT coursemoduleid, userid FROM mdl_course_modules_completion
WHERE coursemoduleid IN ({LECTURES}) AND completionstate IN (1, 2)""",
    "archives": f"""
SELECT DISTINCT contextinstanceid, userid
FROM mdl_logstore_standard_log
WHERE contextlevel = 70 AND contextinstanceid IN ({ARCHIVES})
  AND eventname LIKE '%mod_page%course_module_viewed'""",
    # timecreated подставляется после запроса course
    "logins": """
SELECT userid, COUNT(*) n FROM mdl_logstore_standard_log
WHERE timecreated >= {start} AND eventname LIKE '%core%user_loggedin'
GROUP BY userid""",
}


def unescape(s):
    """Поле из batch-вывода mysql: NULL и экранирование \\\\, \\t, \\n."""
    if s == "NULL":
        return None
    return re.sub(r"\\(.)", lambda m: {"t": "\t", "n": "\n", "0": "\0"}.get(m.group(1), m.group(1)), s)


def query(sql):
    if LOCAL_SQL:
        run = dict(args=[LOCAL_SQL, sql])
    else:
        run = dict(args=SQL_CMD, input=sql)
    out = subprocess.run(**run, capture_output=True, text=True, encoding="utf-8", check=True).stdout
    lines = out.rstrip("\n").split("\n")
    if not lines or not lines[0]:
        return []
    head = lines[0].split("\t")
    return [dict(zip(head, map(unescape, ln.split("\t")))) for ln in lines[1:]]


def fetch():
    data = {k: query(sql) for k, sql in QUERIES.items() if k != "logins"}
    start = int(data["course"][0]["startdate"])
    data["logins"] = query(QUERIES["logins"].format(start=start))
    return data, start


def day(ts):
    return datetime.fromtimestamp(ts, TZ).strftime("%d.%m.%Y")


def build(data, start):
    now = datetime.now(timezone.utc).timestamp()
    zstart = {int(z["id"]): int(z["start_time"]) for z in data["zooms"]}
    lessons = [n for n, (zid, _, _) in LESSONS.items() if zstart.get(zid, now + 1) <= now]

    zoom_lesson = {v[0]: n for n, v in LESSONS.items()}
    archive_lesson = {v[1]: n for n, v in LESSONS.items()}
    lecture_lesson = {v[2]: n for n, v in LESSONS.items()}
    zoom = {(zoom_lesson[int(r["zoomid"])], int(r["userid"])) for r in data["zoom"]}
    archive = {(archive_lesson[int(r["contextinstanceid"])], int(r["userid"])) for r in data["archives"]}
    lecture = {(lecture_lesson[int(r["coursemoduleid"])], int(r["userid"])) for r in data["lectures"]}
    logins = {int(r["userid"]): int(r["n"]) for r in data["logins"]}

    head = ["Имя Фамилия", "Имейл", "Группа", "Ссылка на профиль"]
    notes = [f"Выгрузка из edu.kabacademy.com, курс {COURSE}, на {datetime.now(TZ):%d.%m.%Y %H:%M} (время Израиля).",
             "", "Группа студента в курсе.", ""]
    for n in lessons:
        d = datetime.fromtimestamp(zstart[LESSONS[n][0]], TZ).strftime("%d.%m")
        head += [f"Лекция по уроку {n}, пройдена/нет",
                 f"Вебинар по уроку {n} Zoom ({d}), да/нет",
                 f"Вебинар по уроку {n} архив, да/нет"]
        notes += ["Отметка о выполнении лекции: дочитал до конца.",
                  "Есть в отчёте Zoom. Moodle узнаёт участника по email и имени; "
                  "зашёл в Zoom под чужим именем и почтой — может не попасть.",
                  f"Открывал страницу «Запись вебинара {n} (архив)». "
                  "Сами ролики на YouTube/Drive, досмотрел ли — Moodle не видит."]
    head += [f"Сколько раз заходил на сайт с начала обучения ({day(start)})", "Дата последнего посещения сайта"]
    notes += ["Входы на сайт (событие входа) с даты начала курса. Вход в мобильное приложение не считается.",
              "Последнее действие на сайте, в том числе из мобильного приложения."]

    rows = []
    for u in data["users"]:
        uid = int(u["id"])
        roles = set((u["roles"] or "").split(","))
        if "student" not in roles or roles & STAFF_ROLES or uid in ADMINS:
            continue
        row = [f"{u['firstname']} {u['lastname']}".strip(), u["email"], u["grp"] or "",
               f"{SITE}/user/view.php?id={uid}&course={COURSE}"]
        for n in lessons:
            row += ["пройдена" if (n, uid) in lecture else "нет",
                    "да" if (n, uid) in zoom else "нет",
                    "да" if (n, uid) in archive else "нет"]
        last = int(u["lastaccess"] or 0)
        row += [logins.get(uid, 0),
                datetime.fromtimestamp(last, TZ).strftime("%Y-%m-%d") if last else ""]
        rows.append(row)

    def group_key(g):
        first = g.split(",")[0].strip()
        return (0, int(first), "") if first.isdigit() else (1 if first else 2, 0, first)

    rows.sort(key=lambda r: (group_key(r[2]), r[0].lower()))
    return head, notes, rows, lessons


def save_xlsx(head, notes, rows, outdir, stamp):
    # openpyxl нужен только для копии: на web-18 его нет, там -o не передаётся
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Активность"
    ws.append(head)
    green = PatternFill("solid", fgColor="D9EAD3")
    date_col = len(head) - 1
    for row in rows:
        out = list(row)
        if out[date_col]:
            out[date_col] = datetime.strptime(out[date_col], "%Y-%m-%d")
        ws.append(out)
    for r in ws.iter_rows(min_row=2, min_col=5, max_col=len(head) - 2):
        for c in r:
            if c.value in ("да", "пройдена"):
                c.fill = green
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws.iter_cols(min_col=date_col + 1, max_col=date_col + 1, min_row=2):
        for cell in c:
            cell.number_format = "DD.MM.YYYY"
    ws.row_dimensions[1].height = 60
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    widths = [28, 30, 8, 18] + [12] * (len(head) - 6) + [16, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    fname = Path(outdir) / f"Активность студентов КК1 ({stamp:%Y-%m-%d %H%M}).xlsx"
    wb.save(fname)
    return fname


def post(url, payload):
    """POST в Apps Script. Google отвечает 302 на googleusercontent — urllib сам идёт туда GET-ом."""
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(url, data=body, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=360) as r:
        return json.loads(r.read().decode("utf-8"))


def push(config, head, notes, rows, stamp):
    cfg = json.loads(Path(config).read_text(encoding="utf-8"))
    payload = {
        "secret": cfg["secret"],
        "title": f"{stamp:%d.%m.%Y}",
        "headers": head,
        "notes": notes,
        "rows": rows,
        "textCols": [0, 1],  # имя и почта — с апострофом, чтобы «=…» и «+…» не стали формулой/числом
        "dateCols": [len(head) - 1],
        "markCols": [4, len(head) - 2],  # [с, до) — столбцы уроков, «да»/«пройдена» красятся
    }
    last = None
    for attempt in range(3):  # Google изредка отвечает не тем — повтор безопасен, вкладка заменяется
        try:
            res = post(cfg["url"], payload)
            if res.get("ok") and res.get("rows") == len(rows):
                return res
            last = res
        except Exception as e:  # noqa: BLE001 — сеть/HTML вместо JSON
            last = repr(e)
        time.sleep(15 * (attempt + 1))
    raise SystemExit(f"Не удалось положить выгрузку в таблицу: {last}")


def main():
    global LOCAL_SQL
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--outdir", help="папка для xlsx-копии (нужен openpyxl)")
    ap.add_argument("--push", metavar="CONFIG", help="config.json с url и secret Apps Script")
    ap.add_argument("--local-sql", metavar="SH", help="moodle_sql.sh — SQL на самом web-18, без ssh")
    args = ap.parse_args()
    LOCAL_SQL = args.local_sql
    stamp = datetime.now(TZ)
    print(f"=== {stamp:%d.%m.%Y %H:%M} (Израиль)", flush=True)
    data, start = fetch()
    head, notes, rows, lessons = build(data, start)
    if args.outdir:
        print(save_xlsx(head, notes, rows, args.outdir, stamp))
    print("студентов:", len(rows), "уроков:", lessons)
    for i, n in enumerate(lessons):
        col = 4 + 3 * i
        cnt = [sum(1 for r in rows if r[col + k] != "нет") for k in range(3)]
        print(f"урок {n}: лекция {cnt[0]}, zoom {cnt[1]}, архив {cnt[2]}")
    print("заходили на сайт с начала:", sum(1 for r in rows if r[-2] > 0))
    if args.push:
        res = push(args.push, head, notes, rows, stamp)
        print("таблица:", res)


if __name__ == "__main__":
    main()
